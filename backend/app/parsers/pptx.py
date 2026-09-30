import io
import logging
import os
import uuid
from typing import Any, Dict, List, Optional
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.parsers.base import BaseParser, ExtractedAsset, ExtractedBlock, ExtractedDocument, ExtractedSection

logger = logging.getLogger("accesslearn.parsers.pptx")


class PPTXParser(BaseParser):
    """Parser for Microsoft PowerPoint (.pptx) presentations."""

    def can_parse(self, mime_type: str, filename: str) -> bool:
        ext = os.path.splitext(filename)[1].lower().lstrip(".")
        return "presentationml" in mime_type or mime_type == "application/vnd.ms-powerpoint" or ext in ["pptx", "ppt"]

    def parse(self, file_bytes: bytes, filename: str, options: Optional[Dict[str, Any]] = None) -> ExtractedDocument:
        logger.info(f"Starting PPTX parsing for {filename} ({len(file_bytes)} bytes)")

        try:
            prs = pptx.Presentation(io.BytesIO(file_bytes))
        except Exception as e:
            logger.error(f"Failed to open PPTX document {filename}: {e}")
            raise ValueError(f"Corrupted or invalid PPTX document: {e}")

        doc_title = os.path.splitext(filename)[0].replace("_", " ").title()
        
        # Check first slide title
        if prs.slides and hasattr(prs.slides[0].shapes, "title") and prs.slides[0].shapes.title:
            slide1_title = prs.slides[0].shapes.title.text.strip()
            if slide1_title:
                doc_title = slide1_title

        extracted_doc = ExtractedDocument(
            title=doc_title,
            source_language="en",
            pages_or_slides_count=len(prs.slides),
            metadata={
                "source_type": "pptx",
                "slide_count": len(prs.slides),
                "slide_width": prs.slide_width,
                "slide_height": prs.slide_height,
            }
        )

        all_sections: List[ExtractedSection] = []
        all_assets: List[ExtractedAsset] = []

        for slide_idx, slide in enumerate(prs.slides):
            slide_num = slide_idx + 1
            slide_blocks: List[ExtractedBlock] = []
            slide_title = f"Slide {slide_num}"

            # Extract slide title shape
            if hasattr(slide.shapes, "title") and slide.shapes.title and slide.shapes.title.text.strip():
                slide_title = slide.shapes.title.text.strip()
                slide_blocks.append(
                    ExtractedBlock(
                        type="heading",
                        text=slide_title,
                        page_or_slide=slide_num,
                        position=1,
                        metadata={"heading_level": 2, "source_type": "pptx", "slide_number": slide_num}
                    )
                )

            # Extract speaker notes if present
            notes_text = None
            try:
                if slide.has_notes_slide and slide.notes_slide:
                    notes_tf = slide.notes_slide.notes_text_frame
                    if notes_tf and notes_tf.text.strip():
                        notes_text = notes_tf.text.strip()
            except Exception as notes_err:
                logger.warning(f"Could not extract notes on slide {slide_num}: {notes_err}")

            # Iterate through slide shapes
            position = len(slide_blocks) + 1
            for shape in slide.shapes:
                # Skip title shape already processed
                if hasattr(slide.shapes, "title") and shape == slide.shapes.title:
                    continue

                # 1. Text Frame
                if shape.has_text_frame and shape.text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if not text:
                            continue

                        # Check level for list bullets
                        if paragraph.level > 0 or text.startswith(("•", "-", "*", "–")):
                            slide_blocks.append(
                                ExtractedBlock(
                                    type="list",
                                    text=text,
                                    page_or_slide=slide_num,
                                    position=position,
                                    metadata={"level": paragraph.level, "source_type": "pptx", "slide_number": slide_num}
                                )
                            )
                        else:
                            slide_blocks.append(
                                ExtractedBlock(
                                    type="paragraph",
                                    text=text,
                                    page_or_slide=slide_num,
                                    position=position,
                                    metadata={"level": paragraph.level, "source_type": "pptx", "slide_number": slide_num}
                                )
                            )
                        position += 1

                # 2. Table
                elif shape.has_table and shape.table:
                    matrix = []
                    for row in shape.table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells]
                        matrix.append(row_cells)

                    if matrix:
                        headers = matrix[0] if len(matrix) > 1 else []
                        rows = matrix[1:] if len(matrix) > 1 else matrix
                        slide_blocks.append(
                            ExtractedBlock(
                                type="table",
                                page_or_slide=slide_num,
                                position=position,
                                metadata={
                                    "headers": headers,
                                    "rows": rows,
                                    "row_count": len(matrix),
                                    "col_count": len(matrix[0]) if matrix else 0,
                                    "source_type": "pptx",
                                    "slide_number": slide_num
                                }
                            )
                        )
                        position += 1

                # 3. Picture / Image
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        img = shape.image
                        img_bytes = img.blob
                        img_ext = img.ext or "png"
                        asset_id = str(uuid.uuid4())
                        asset = ExtractedAsset(
                            asset_id=asset_id,
                            type="image",
                            filename=f"slide_{slide_num}_img_{len(all_assets) + 1}.{img_ext}",
                            mime_type=f"image/{img_ext}",
                            data_bytes=img_bytes,
                            page_or_slide=slide_num,
                            metadata={"slide_number": slide_num}
                        )
                        all_assets.append(asset)
                        slide_blocks.append(
                            ExtractedBlock(
                                type="image",
                                text=f"[Slide {slide_num} Image]",
                                page_or_slide=slide_num,
                                position=position,
                                asset_id=asset_id,
                                metadata={"asset_filename": asset.filename, "source_type": "pptx", "slide_number": slide_num}
                            )
                        )
                        position += 1
                    except Exception as img_err:
                        logger.warning(f"Failed to extract picture on slide {slide_num}: {img_err}")
                        extracted_doc.warnings.append(f"Slide {slide_num}: image extraction failed.")

            # Append speaker notes as a media/notes block if present
            if notes_text:
                slide_blocks.append(
                    ExtractedBlock(
                        type="paragraph",
                        text=f"Speaker Notes: {notes_text}",
                        page_or_slide=slide_num,
                        position=position,
                        metadata={"is_speaker_notes": True, "source_type": "pptx", "slide_number": slide_num}
                    )
                )

            all_sections.append(
                ExtractedSection(
                    title=f"Slide {slide_num}: {slide_title}" if slide_title != f"Slide {slide_num}" else slide_title,
                    blocks=slide_blocks,
                    page_or_slide=slide_num
                )
            )

        extracted_doc.sections = all_sections
        extracted_doc.assets = all_assets
        logger.info(f"PPTX parsed: {len(all_sections)} slides, {len(all_assets)} assets")
        return extracted_doc
